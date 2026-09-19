"""解析内部系统导出的 xlsx —— 唯一读那份文件的地方。

契约：xlsx 的表头就是 Base 看板表的列，逐字对应 `schema.DAILY_BOARD_FIELDS`，顺序也
一样。少任何一列直接拒绝（缺列＝下游拿空值往下算，宁可在导入前停下）；多出来的列忽略
并提示。导出工具偶尔把括号写成全角、表头前后带空格，这类差异规整后再比。

这个模块**只做解析**：不认识 Base、不认识网络、不打印给人看的东西。输出的 `BoardRow`
是纯数据。
"""

from __future__ import annotations

import logging
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from ..domain import schema
from ..lark.bitable import FIELD_TYPE_DATETIME, FIELD_TYPE_NUMBER
from ..lark.values import PrecisionLossError, extract_text, to_uid

logger = logging.getLogger(__name__)

# xlsx 必须有的表头，也就是 Base 里看板表的列。
EXPECTED_HEADERS: tuple[str, ...] = tuple(schema.DAILY_BOARD_FIELDS)

DATE_FIELDS = frozenset(
    name for name, kind in schema.DAILY_BOARD_FIELDS.items() if kind == FIELD_TYPE_DATETIME
)
NUMBER_FIELDS = frozenset(
    name for name, kind in schema.DAILY_BOARD_FIELDS.items() if kind == FIELD_TYPE_NUMBER
)

# 一行里缺了这几个值就没法算佣金，跳过这一行。其余列都允许为空。
REQUIRED_VALUES: tuple[str, ...] = (
    schema.BOARD_CLIENT_UID,
    schema.BOARD_ORDER_DATE,
    schema.BOARD_TOTAL_REVENUE,
)

# 导出工具偶尔写成全角的符号，规整成 schema 里用的半角再比。
_HEADER_TRANSLATION = str.maketrans({"（": "(", "）": ")", "＿": "_", "＋": "+"})
_WHITESPACE = re.compile(r"\s+")


def normalize_header(raw: Any) -> str:
    """表头规整：去掉所有空白，全角括号、下划线、加号换成半角。"""
    if raw is None:
        return ""
    return _WHITESPACE.sub("", str(raw)).translate(_HEADER_TRANSLATION)


@dataclass(frozen=True)
class BoardRow:
    order_date: date
    fields: dict[str, Any]


class BoardImportError(RuntimeError):
    """xlsx 结构或值不符合要求。"""


def _load_workbook(path: Path):
    """惰性 import openpyxl，避免让主进程启动就吃这个依赖。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BoardImportError("缺 openpyxl，先跑 `uv sync` 装齐依赖再试。") from exc

    if not path.exists():
        raise BoardImportError(f"找不到文件：{path}")
    # data_only=True 让公式单元格返回上次计算的值，而不是公式本身
    return load_workbook(path, data_only=True, read_only=True)


def _read_header(sheet) -> list[str]:
    """读第一行表头，返回每一列对应的看板列名，认不出的列是空串。

    缺任何一列都报错：少了哪列，下游就会拿着空值往下算，宁可在导入前停下。
    同一列出现两次也报错：不知道该用哪一列。
    """
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if first_row is None:
        raise BoardImportError("xlsx 是空的，连表头都没有")

    raw_headers = ["" if cell is None else str(cell) for cell in first_row]
    expected = set(EXPECTED_HEADERS)
    normalized = [normalize_header(raw) for raw in raw_headers]
    columns = [name if name in expected else "" for name in normalized]

    missing = [name for name in EXPECTED_HEADERS if name not in columns]
    if missing:
        raise BoardImportError(
            f"xlsx 表头缺少这些列：{'、'.join(missing)}\n"
            f"实际表头：{[raw for raw in raw_headers if raw]}\n"
            "表头以 2026-09-17 的「OTC组销售明细」导出为准，清单见 schema.DAILY_BOARD_FIELDS。"
        )

    duplicated = sorted({name for name in columns if name and columns.count(name) > 1})
    if duplicated:
        raise BoardImportError(f"xlsx 表头里这些列出现了不止一次：{'、'.join(duplicated)}")

    unknown = [
        raw for raw, name in zip(raw_headers, columns, strict=True) if raw.strip() and not name
    ]
    if unknown:
        logger.warning("忽略看板表里没有的列：%s", unknown)

    return columns


def _cell_to_uid(value: Any, *, row_num: int) -> str:
    """用户ID 的值。任何浮点形态都当作精度已损伤，直接拒绝。

    xlsx 里用户ID 应该是文本；拿到 float 就是源头把它当数字存了，18-19 位的数字
    过一遍 float 就丢低位，下游对不上。让这里就炸，而不是继续往下算。
    """
    column = schema.BOARD_CLIENT_UID
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        raise BoardImportError(f"第 {row_num} 行的「{column}」是布尔值")
    if isinstance(value, float):
        raise BoardImportError(
            f"第 {row_num} 行的「{column}」是浮点数，说明源头把这一列当数字存了。"
            f"请让上游把「{column}」那一列在 xlsx 里设成文本格式重新导出，"
            "否则 18-19 位的 ID 低位已经被抹成 0。"
        )
    if isinstance(value, int):
        return str(value)
    try:
        return to_uid(value)
    except PrecisionLossError as exc:
        raise BoardImportError(f"第 {row_num} 行的「{column}」无法安全转成字符串：{exc}") from exc


def _cell_to_date(value: Any, column: str, *, row_num: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        # 只取日期部分（扔掉可能跟着的时分秒），再规范化分隔符
        text = value.strip().split(" ", 1)[0].replace("/", "-")
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            pass
    raise BoardImportError(f"第 {row_num} 行的「{column}」无法解析成日期：{value!r}")


def _cell_to_number(value: Any, column: str, *, row_num: int) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):  # bool 是 int 的子类，得先挡
        raise BoardImportError(f"第 {row_num} 行的「{column}」是布尔值 {value!r}")
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, str):
        stripped = value.strip().replace(",", "")
        if not stripped:
            return None
        try:
            return float(stripped)
        except ValueError as exc:
            raise BoardImportError(f"第 {row_num} 行的「{column}」不是数字：{value!r}") from exc
    raise BoardImportError(f"第 {row_num} 行的「{column}」无法转数字：{value!r}")


def _cell_value(cell: Any, column: str, *, row_num: int) -> Any:
    """按列的类型转换一个单元格。返回 None 表示这个单元格是空的，不写这一列。"""
    if column == schema.BOARD_CLIENT_UID:
        return _cell_to_uid(cell, row_num=row_num) or None
    if column in DATE_FIELDS:
        return _cell_to_date(cell, column, row_num=row_num)
    if column in NUMBER_FIELDS:
        return _cell_to_number(cell, column, row_num=row_num)
    return extract_text(cell) or None


def parse_workbook(path: Path, *, only_date: date | None = None) -> list[BoardRow]:
    """把 xlsx 读成 BoardRow 列表。所有校验都在这里做，不碰 Base。

    ``only_date`` 传了就只保留那一天的行。
    """
    workbook = _load_workbook(path)
    rows: list[BoardRow] = []
    skipped = 0
    try:
        sheet = workbook.active
        columns = _read_header(sheet)

        for index, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None or cell == "" for cell in raw_row):
                continue  # 空行

            fields: dict[str, Any] = {}
            for cell, column in zip(raw_row, columns, strict=False):
                if not column or cell is None or cell == "":
                    continue
                value = _cell_value(cell, column, row_num=index)
                if value is not None:
                    fields[column] = value

            # 缺值按「这一列有没有值」判断，不按真假判断：总收入是 0 的行照样要导。
            missing = [column for column in REQUIRED_VALUES if column not in fields]
            if missing:
                # 只报行号和缺了哪列。日志会被转发和截图，客户名称、金额这些不打出来。
                labels = "、".join(f"「{column}」" for column in missing)
                logger.warning("第 %d 行缺少%s，跳过", index, labels)
                skipped += 1
                continue

            order_date: date = fields[schema.BOARD_ORDER_DATE]
            if only_date is not None and order_date != only_date:
                continue

            rows.append(BoardRow(order_date=order_date, fields=fields))
    finally:
        workbook.close()

    if skipped:
        logger.warning("一共跳过 %d 行缺少必填值的记录", skipped)
    return rows


def split_by_station(rows: list[BoardRow]) -> tuple[list[BoardRow], Counter[str]]:
    """只留看板要的站点，顺带数出每个站点各有多少行，方便核对筛得对不对。

    这是**口径的执行点**：哪些站点进 Base 由 `schema.BOARD_STATION_IN_SCOPE` 决定，
    不在这里写死站点名。
    """
    counts: Counter[str] = Counter(
        row.fields.get(schema.BOARD_STATION, "（站点为空）") for row in rows
    )
    kept = [
        row for row in rows if row.fields.get(schema.BOARD_STATION) == schema.BOARD_STATION_IN_SCOPE
    ]
    return kept, counts


def describe_stations(counts: Counter[str]) -> str:
    """一行话说明筛完剩下什么：「新加坡站 1650 行导入；香港站 10640 行不导入」。"""
    in_scope = schema.BOARD_STATION_IN_SCOPE
    text = f"{in_scope} {counts.get(in_scope, 0)} 行导入"
    others = [f"{station} {n} 行" for station, n in counts.most_common() if station != in_scope]
    if others:
        text += f"；{'、'.join(others)}不导入"
    return text
