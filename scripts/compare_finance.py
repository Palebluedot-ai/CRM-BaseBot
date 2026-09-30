#!/usr/bin/env python
"""拿财务的月度 Referral Fee 表（xlsx）对我们的结算，逐渠道、逐客户说出差在哪。**只读。**

    uv run python scripts/compare_finance.py --period 2026-07 \
        --file ~/Downloads/JulyReferralFee_Recomputed.xlsx

读财务表的「Summary」页：有「Referrer Code」的是渠道行，下面没有编号、有「Client Name」的是
它的客户行，「GRAND TOTAL」是合计。和三样东西比：

  · 结算表（Commission Summary）里这个月每个渠道的应付 —— 实际付的钱
  · 按**现在**的资料重算的每个渠道、每个客户（和看板、仪表盘明细同一套算法）
  · 客户名字对不上时，说是我们没这个客户（没登记 / UID 对不上），还是财务没有

差额在 1 块钱以内的当作四舍五入（财务是逐笔先取两位再加，我们是加完再取），不列出来。
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

ZERO = Decimal("0")
ROUNDING = Decimal("1.00")


def norm(name: str) -> str:
    """名字归一：大写、只留字母数字。「CO., LIMITED」和「CO LIMITED」算同一个。"""
    return re.sub(r"[^0-9A-Z一-鿿]", "", (name or "").upper())


def _dec(value) -> Decimal:
    if value in (None, ""):
        return ZERO
    return Decimal(str(value)).quantize(Decimal("0.01"))


@dataclass
class Channel:
    code: str
    name: str
    rate: str
    revenue: Decimal
    commission: Decimal
    clients: dict[str, tuple[str, Decimal, Decimal]] = field(default_factory=dict)
    """norm(客户名) -> (客户名, 收入, 佣金)"""


@dataclass
class FinanceMonth:
    channels: dict[str, Channel]
    grand_total: Decimal


def read_finance(path: Path) -> FinanceMonth:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Summary"] if "Summary" in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(v or "").strip() for v in next(rows)]

    def col(*names: str) -> int:
        for name in names:
            if name in header:
                return header.index(name)
        raise SystemExit(f"财务表的 Summary 页找不到这一列：{names[0]}（表头是 {header}）")

    c_code, c_name = col("Referrer Code"), col("Name of Referrer")
    c_rate, c_client = col("Commission Rate (%)"), col("Client Name")
    c_rev = col("Total PnL (USD)", "Total Revenue (USD)")
    c_fee = col("Total Commission (USD)")

    channels: dict[str, Channel] = {}
    grand = ZERO
    current: Channel | None = None
    for row in rows:
        code = str(row[c_code] or "").strip()
        name = str(row[c_name] or "").strip()
        client = str(row[c_client] or "").strip()
        if name.upper() == "GRAND TOTAL":
            grand = _dec(row[c_fee])
            continue
        if code:
            current = Channel(
                code, name, str(row[c_rate] or ""), _dec(row[c_rev]), _dec(row[c_fee])
            )
            channels[code] = current
        elif client and current is not None:
            current.clients[norm(client)] = (client, _dec(row[c_rev]), _dec(row[c_fee]))
    return FinanceMonth(channels, grand)


def compare(
    finance: FinanceMonth,
    settled: dict[str, tuple[str, Decimal]],
    now: dict[str, Channel],
) -> list[str]:
    """返回要给人看的行。

    ``settled`` 编号 -> (名称, 结算应付)；``now`` 是现算（客户同样按 norm 名字）。
    """
    lines: list[str] = []
    codes = sorted(set(finance.channels) | set(settled) | set(now))
    for code in codes:
        fin = finance.channels.get(code)
        paid = settled.get(code)
        live = now.get(code)
        fin_fee = fin.commission if fin else ZERO
        paid_fee = paid[1] if paid else ZERO
        if abs(fin_fee - paid_fee) <= ROUNDING:
            continue
        name = (fin.name if fin else "") or (paid[0] if paid else "") or (live.name if live else "")
        lines.append(
            f"{code} {name}：财务 {fin_fee:,.2f}　结算 {paid_fee:,.2f}"
            f"　差 {paid_fee - fin_fee:+,.2f}"
        )
        if live is not None and abs(live.commission - paid_fee) > ROUNDING:
            lines.append(
                f"    · 按现在的资料重算是 {live.commission:,.2f}（结算后改过比例或补登记了客户）"
            )
        fin_clients = fin.clients if fin else {}
        live_clients = live.clients if live else {}
        for key in sorted(set(fin_clients) | set(live_clients)):
            f = fin_clients.get(key)
            w = live_clients.get(key)
            if f and not w:
                lines.append(
                    f"    · 财务有、我们没有：{f[0]}（收入 {f[1]:,.2f}，佣金 {f[2]:,.2f}）"
                    " —— 客户没登记，或者登记的 UID 和看板对不上"
                )
            elif w and not f:
                lines.append(
                    f"    · 我们有、财务没有：{w[0]}（收入 {w[1]:,.2f}，佣金 {w[2]:,.2f}）"
                )
            elif f and w and abs(f[1] - w[1]) > ROUNDING:
                lines.append(f"    · {f[0]}：收入 财务 {f[1]:,.2f} / 我们 {w[1]:,.2f}")
        if (
            fin
            and live
            and fin.rate
            and live.rate
            and fin.rate.rstrip("%") != live.rate.rstrip("%")
        ):
            lines.append(f"    · 比例：财务 {fin.rate} / 我们现在 {live.rate}")
    return lines


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="财务的月度佣金表 vs 我们的结算")
    parser.add_argument("--file", required=True, help="财务那份 xlsx")
    parser.add_argument("--period", required=True, help="YYYY-MM")
    args = parser.parse_args(argv)

    from crm_basebot.domain.commission_query import CommissionQueryService
    from crm_basebot.jobs.archive import _EVERYONE, KIND_TRADE, load_settled
    from crm_basebot.lark.bitable import BitableClient
    from crm_basebot.startup import load_settings, require_settings

    finance = read_finance(Path(args.file).expanduser())
    settings = load_settings()
    require_settings(settings, "LARK_BASE_APP_TOKEN", "TABLE_COMMISSION", "TABLE_DAILY_BOARD")
    bitable = BitableClient(settings.base_app_token)

    settled = {
        s.referral_no: (s.name, s.amount)
        for s in load_settled(bitable, settings, args.period)
        if s.kind == KIND_TRADE
    }
    result = CommissionQueryService(bitable, settings=settings).query(_EVERYONE, [args.period])
    now: dict[str, Channel] = {}
    for referral in result.referrals_in(args.period):
        shares = referral.client_shares()
        channel = Channel(
            referral.referral_no,
            referral.referral_name,
            f"{referral.rate_percent.normalize():f}%",
            referral.revenue_total,
            referral.payable,
        )
        for uid, client in referral.clients.items():
            channel.clients[norm(client.name or uid)] = (
                client.name or uid,
                client.revenue.quantize(Decimal("0.01")),
                shares.get(uid, ZERO),
            )
        now[referral.referral_no] = channel

    paid_total = sum((fee for _, fee in settled.values()), ZERO)
    print("只读，没有改任何东西。\n")
    print(
        f"{args.period} 交易佣金：财务 {finance.grand_total:,.2f}　结算 {paid_total:,.2f}"
        f"　差 {paid_total - finance.grand_total:+,.2f}\n"
    )
    lines = compare(finance, settled, now)
    if not lines:
        print("每个渠道都对得上（1 块钱以内算四舍五入）。")
    else:
        print("\n".join(lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
