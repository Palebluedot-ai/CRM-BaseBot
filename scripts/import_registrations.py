#!/usr/bin/env python
"""把渠道登记表和渠道客户表从模板 xlsx 导进 Base，重复跑只更新不重复。

模板是 2026-09-17 给的「Template .xlsx」，三个工作表：

    Referral Registration   渠道：Referral Code、Name、Email、Start Date、Commission Rate、
                            Payout Frequency、Submitted On、Sales In Charge
    Referred Clients        客户：Referral Code、Client Name、UID、Sales In Charge
    用户UID                  客户全量 UID 和名称。客户行的 UID 空着时，按客户名从这里补

列名和 Base 里的对应关系见 schema.py 表 1、表 2 的注释。

## 重复跑

渠道按编号找旧行，客户先按 UID 找、找不到再按「编号 + 客户名」找；找到就比对每个字段，
有变化才更新，没有就跳过。所以模板改了再跑一次就能同步，不会堆出重复行。
客户行原来没有 UID、后来在模板里补上了，也是更新同一行。

## UID 只认可靠的

18-19 位的 UID 经不起 Excel：存成数字会被抹掉低位，读回来是浮点数；用文本形式贴过去
又常常已经是被抹过的（尾巴一串 0）或科学计数法。所以：客户行的 UID 是浮点数直接报错
停下；用户UID 表里尾零、科学计数法形态的值不用，留空并提示。留空的客户导进去也没坏处，
只是算不到佣金，等补上 UID 再跑一次。

## 归属人先空着

归属销售、登记人OpenID 两列这次不写：模板里只有销售的姓名，没有 open_id。姓名写进
「负责销售」列，销售名册也按姓名补齐、OpenID 留空。拿到 open_id 之后填进名册，再按姓名
回填归属。

## 用法

    uv run python scripts/import_registrations.py --file "Template .xlsx"          # 预演
    uv run python scripts/import_registrations.py --file "Template .xlsx" --apply  # 真的写
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, tzinfo
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from crm_basebot.domain import schema  # noqa: E402
from crm_basebot.domain.dates import date_to_ms  # noqa: E402
from crm_basebot.domain.referral import display_title  # noqa: E402
from crm_basebot.lark.bitable import BitableClient, Record  # noqa: E402
from crm_basebot.lark.values import (  # noqa: E402
    PrecisionLossError,
    extract_text,
    looks_excel_truncated,
    looks_scientific_notation,
    to_number,
    to_uid,
)
from crm_basebot.startup import load_settings, require_settings  # noqa: E402

logger = logging.getLogger(__name__)

SHEET_REFERRALS = "Referral Registration"
SHEET_CLIENTS = "Referred Clients"
SHEET_UIDS = "用户UID"

# 模板表头 -> 内部键。按规整后的表头（去空白、小写）认；邮箱那一列单独处理，
# 因为模板里它的表头被人覆盖成了一个邮箱地址。
REFERRAL_COLUMNS = {
    "referral code": "code",
    "referrer code": "code",
    "name": "name",
    "start date": "start_date",
    "commission rate": "rate",
    "payout frequency": "payout",
    "submitted on": "submitted_on",
    "sales in charge": "sales",
}
REFERRAL_REQUIRED = (
    "code",
    "name",
    "email",
    "start_date",
    "rate",
    "payout",
    "submitted_on",
    "sales",
)
CLIENT_COLUMNS = {
    "referral code": "code",
    "referrer code": "code",
    "client name": "name",
    "uid": "uid",
    "sales in charge": "sales",
}
CLIENT_REQUIRED = ("code", "name")
UID_COLUMNS = {"uid": "uid", "客户名称": "name", "client name": "name"}


class RegistrationImportError(RuntimeError):
    """模板结构或值不符合要求。"""


@dataclass(frozen=True)
class ReferralRow:
    code: str
    name: str
    email: str
    start_date: date | None
    rate_percent: float | None
    payout: str
    submitted_on: date | None
    sales: str
    status: str


@dataclass(frozen=True)
class ClientRow:
    code: str
    name: str
    uid: str
    sales: str


@dataclass
class ParsedTemplate:
    referrals: list[ReferralRow]
    clients: list[ClientRow]
    uid_by_name: dict[str, str]
    warnings: list[str] = field(default_factory=list)


@dataclass
class Summary:
    referrals_created: int = 0
    referrals_updated: int = 0
    clients_created: int = 0
    clients_updated: int = 0
    sales_created: int = 0
    notes: list[str] = field(default_factory=list)


# ---------- 规整 ----------


def _norm_header(raw: Any) -> str:
    return re.sub(r"\s+", " ", str(raw or "")).strip().casefold()


def _norm_name(raw: Any) -> str:
    """客户名做键：大写、空白压成一个。模板里两张表的写法可能只差大小写和空格。"""
    return re.sub(r"\s+", " ", str(raw or "")).strip().upper()


def _clean_text(raw: Any) -> str:
    """名字、邮箱这类值只去首尾空白，中间照原样，别替人改数据。"""
    return extract_text(raw).strip()


def _clean_sales(raw: Any) -> str:
    """模板里销售是 @ 提及：去掉开头的 @，空白压成一个，好和名册按姓名对上。"""
    return re.sub(r"\s+", " ", extract_text(raw)).lstrip("@").strip()


def _clean_code(raw: Any) -> str:
    return _clean_text(raw).upper()


def _cell_to_date(value: Any, column: str, *, row_num: int, warnings: list[str]) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().split(" ", 1)[0].replace("/", "-")
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        warnings.append(f"第 {row_num} 行的「{column}」看不出是日期，留空")
        return None


def _rate_percent(value: Any) -> float | None:
    """分佣比例统一成百分数。模板里写的是小数（0.2 = 20%），也认「20%」和 20。"""
    if value in (None, ""):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        explicit_percent = text.endswith("%")
        text = text.rstrip("%").strip()
        if not text:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise RegistrationImportError(f"分佣比例不是数字：{value!r}") from exc
        if not explicit_percent and 0 <= number <= 1:
            number *= 100
    else:
        number = Decimal(str(value))
        if 0 <= number <= 1:
            number *= 100
    return float(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _status_for(name: str) -> str:
    """渠道名里带 TERMINATED 的按停用导入。佣金计算不看状态，这只影响人看。"""
    return schema.STATUS_DISABLED if "TERMINAT" in name.upper() else schema.STATUS_ACTIVE


def _cell_to_uid(value: Any, *, where: str) -> str:
    """UID 单元格。浮点数说明精度已损，直接报错；文本形态先去掉引号再交给 to_uid。"""
    if value in (None, ""):
        return ""
    if isinstance(value, bool):
        raise RegistrationImportError(f"{where}的 UID 是布尔值")
    if isinstance(value, float):
        raise RegistrationImportError(
            f"{where}的 UID 是浮点数，说明这个单元格被当成数字存了，18-19 位的低位已经"
            "被抹掉。请把那一格设成文本后重新粘贴。"
        )
    if isinstance(value, int):
        return str(value)
    text = str(value).strip().strip("'‘’\"“”").strip()
    try:
        return to_uid(text)
    except PrecisionLossError as exc:
        raise RegistrationImportError(f"{where}的 UID 无法安全转成字符串：{exc}") from exc


def _damaged_uid_reason(uid: str) -> str | None:
    if looks_scientific_notation(uid):
        return "是科学计数法"
    if looks_excel_truncated(uid):
        return "尾巴是一串 0，像被 Excel 抹过低位"
    if not uid.isdigit():
        return "不是纯数字"
    return None


# ---------- 读模板 ----------


def _find_sheet(workbook, title: str):
    if title in workbook.sheetnames:
        return workbook[title]
    wanted = _norm_header(title)
    for name in workbook.sheetnames:
        if wanted in _norm_header(name):
            return workbook[name]
    raise RegistrationImportError(f"模板里没有工作表「{title}」，只有 {workbook.sheetnames}")


def _rows(sheet) -> list[tuple]:
    return [
        row
        for row in sheet.iter_rows(values_only=True)
        if any(cell not in (None, "") for cell in row)
    ]


def _columns(
    header: tuple, mapping: dict[str, str], required: tuple[str, ...], *, sheet: str
) -> dict[str, int]:
    """表头 -> 列下标。邮箱列按「含 email 或 @」认。"""
    found: dict[str, int] = {}
    for index, raw in enumerate(header):
        key = _norm_header(raw)
        if not key:
            continue
        if key in mapping:
            found.setdefault(mapping[key], index)
        elif "email" in key or "@" in key or key == "邮箱":
            found.setdefault("email", index)
    missing = [key for key in required if key not in found]
    if missing:
        raise RegistrationImportError(
            f"工作表「{sheet}」缺少这些列：{missing}；实际表头：{[h for h in header if h]}"
        )
    return found


def _parse_referrals(sheet, warnings: list[str]) -> list[ReferralRow]:
    rows = _rows(sheet)
    if not rows:
        raise RegistrationImportError(f"工作表「{SHEET_REFERRALS}」是空的")
    col = _columns(rows[0], REFERRAL_COLUMNS, REFERRAL_REQUIRED, sheet=SHEET_REFERRALS)

    def cell(row: tuple, key: str) -> Any:
        index = col[key]
        return row[index] if index < len(row) else None

    referrals: list[ReferralRow] = []
    for row_num, row in enumerate(rows[1:], start=2):
        code = _clean_code(cell(row, "code"))
        name = _clean_text(cell(row, "name"))
        if not code or not name:
            warnings.append(f"「{SHEET_REFERRALS}」第 {row_num} 行缺编号或名称，跳过")
            continue
        referrals.append(
            ReferralRow(
                code=code,
                name=name,
                email=_clean_text(cell(row, "email")),
                start_date=_cell_to_date(
                    cell(row, "start_date"), "Start Date", row_num=row_num, warnings=warnings
                ),
                rate_percent=_rate_percent(cell(row, "rate")),
                payout=_clean_text(cell(row, "payout")),
                submitted_on=_cell_to_date(
                    cell(row, "submitted_on"), "Submitted On", row_num=row_num, warnings=warnings
                ),
                sales=_clean_sales(cell(row, "sales")),
                status=_status_for(name),
            )
        )
    return referrals


def _parse_uids(sheet, warnings: list[str]) -> dict[str, str]:
    """用户UID 表 -> {客户名: UID}。坏掉的 UID 不收，同名多个 UID 只收第一个。"""
    rows = _rows(sheet)
    if len(rows) < 2:
        return {}
    col = _columns(rows[0], UID_COLUMNS, ("uid", "name"), sheet=SHEET_UIDS)
    mapping: dict[str, str] = {}
    for row_num, row in enumerate(rows[1:], start=2):
        name = _norm_name(row[col["name"]] if col["name"] < len(row) else None)
        raw = row[col["uid"]] if col["uid"] < len(row) else None
        if not name or raw in (None, ""):
            continue
        uid = _cell_to_uid(raw, where=f"「{SHEET_UIDS}」第 {row_num} 行")
        reason = _damaged_uid_reason(uid)
        if reason:
            warnings.append(f"「{SHEET_UIDS}」第 {row_num} 行的 UID {uid} {reason}，没有采用")
            continue
        if name in mapping and mapping[name] != uid:
            warnings.append(f"「{SHEET_UIDS}」里客户「{name}」对应不止一个 UID，只用了第一个")
            continue
        mapping[name] = uid
    return mapping


def _parse_clients(sheet, uid_by_name: dict[str, str], warnings: list[str]) -> list[ClientRow]:
    rows = _rows(sheet)
    if not rows:
        raise RegistrationImportError(f"工作表「{SHEET_CLIENTS}」是空的")
    col = _columns(rows[0], CLIENT_COLUMNS, CLIENT_REQUIRED, sheet=SHEET_CLIENTS)

    def cell(row: tuple, key: str) -> Any:
        index = col.get(key)
        return row[index] if index is not None and index < len(row) else None

    clients: list[ClientRow] = []
    for row_num, row in enumerate(rows[1:], start=2):
        code = _clean_code(cell(row, "code"))
        name = _clean_text(cell(row, "name"))
        if not code or not name:
            warnings.append(f"「{SHEET_CLIENTS}」第 {row_num} 行缺编号或客户名，跳过")
            continue
        uid = _cell_to_uid(cell(row, "uid"), where=f"「{SHEET_CLIENTS}」第 {row_num} 行")
        if uid:
            reason = _damaged_uid_reason(uid)
            if reason:
                warnings.append(f"「{SHEET_CLIENTS}」第 {row_num} 行的 UID {uid} {reason}，留空")
                uid = ""
        if not uid:
            uid = uid_by_name.get(_norm_name(name), "")
        clients.append(
            ClientRow(code=code, name=name, uid=uid, sales=_clean_sales(cell(row, "sales")))
        )
    return clients


def parse_template(path: Path) -> ParsedTemplate:
    """把模板读成结构化的行。所有校验都在这里，不碰 Base。"""
    from openpyxl import load_workbook

    if not Path(path).exists():
        raise RegistrationImportError(f"找不到文件：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    warnings: list[str] = []
    try:
        referrals = _parse_referrals(_find_sheet(workbook, SHEET_REFERRALS), warnings)
        uid_by_name = _parse_uids(_find_sheet(workbook, SHEET_UIDS), warnings)
        clients = _parse_clients(_find_sheet(workbook, SHEET_CLIENTS), uid_by_name, warnings)
    finally:
        workbook.close()
    return ParsedTemplate(
        referrals=referrals, clients=clients, uid_by_name=uid_by_name, warnings=warnings
    )


# ---------- 和 Base 比对、写入 ----------


def _link_ids(value: Any) -> list[str]:
    if isinstance(value, list):
        return [
            item if isinstance(item, str) else str(item.get("record_id") or item.get("id") or "")
            for item in value
        ]
    if isinstance(value, dict):
        return list(value.get("link_record_ids") or [])
    return []


def _same(existing: Any, desired: Any) -> bool:
    """Base 读回来的值和我们要写的值是不是一回事。读回来的文本是富文本数组，数字可能是 float。"""
    if isinstance(desired, list):
        return _link_ids(existing) == desired
    if isinstance(desired, bool):
        return existing == desired
    if isinstance(desired, int | float):
        number = to_number(existing)
        return number is not None and number == float(desired)
    return extract_text(existing) == str(desired)


def _changes(existing: dict[str, Any], desired: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in desired.items() if not _same(existing.get(key), value)}


def _referral_fields(row: ReferralRow, primary: str, tz: tzinfo) -> dict[str, Any]:
    fields: dict[str, Any] = {
        schema.REFERRAL_NO: row.code,
        schema.REFERRAL_NAME: row.name,
        schema.REFERRAL_EMAIL: row.email,
        schema.REFERRAL_STATUS: row.status,
    }
    if row.start_date is not None:
        fields[schema.REFERRAL_START_DATE] = date_to_ms(row.start_date, tz=tz)
    if row.rate_percent is not None:
        fields[schema.REFERRAL_RATE] = row.rate_percent
    if row.payout:
        fields[schema.REFERRAL_PAYOUT] = row.payout
    if row.submitted_on is not None:
        fields[schema.REFERRAL_SUBMITTED_ON] = date_to_ms(row.submitted_on, tz=tz)
    if row.sales:
        fields[schema.REFERRAL_SALES_NAME] = row.sales
    if primary not in fields:
        # 关联字段展示的是主字段，写成「R001 名称」才不会显示成「无标题记录」
        fields[primary] = display_title(row.code, row.name)
    return fields


def _client_fields(row: ClientRow, referral_record_id: str, primary: str) -> dict[str, Any]:
    fields: dict[str, Any] = {
        schema.CLIENT_NAME: row.name,
        schema.CLIENT_REFERRAL_LINK: [referral_record_id],
    }
    if row.uid:
        fields[schema.CLIENT_UID] = row.uid
    if row.sales:
        fields[schema.CLIENT_SALES_NAME] = row.sales
    if primary not in fields:
        fields[primary] = row.name
    return fields


def _dedupe_clients(clients: Iterable[ClientRow], notes: list[str]) -> list[ClientRow]:
    """同一渠道下同一个客户名只留一行；同一个 UID 出现在两个渠道下只认先出现的。"""
    seen_key: set[tuple[str, str]] = set()
    seen_uid: dict[str, str] = {}
    kept: list[ClientRow] = []
    for row in clients:
        key = (row.code, _norm_name(row.name))
        if key in seen_key:
            notes.append(
                f"渠道 {row.code} 下的客户「{row.name}」在模板里出现了不止一次，只保留一行"
            )
            continue
        if row.uid and row.uid in seen_uid and seen_uid[row.uid] != row.code:
            notes.append(
                f"UID {row.uid} 同时挂在渠道 {seen_uid[row.uid]} 和 {row.code} 下，"
                f"只认先出现的 {seen_uid[row.uid]}，「{row.name}」这一行没导"
            )
            continue
        seen_key.add(key)
        if row.uid:
            seen_uid.setdefault(row.uid, row.code)
        kept.append(row)
    return kept


def _sales_names(parsed: ParsedTemplate) -> list[str]:
    """两张表里出现过的销售姓名，按不区分大小写去重，保留先出现的写法。"""
    names: dict[str, str] = {}
    for name in [r.sales for r in parsed.referrals] + [c.sales for c in parsed.clients]:
        if name and name.casefold() not in names:
            names[name.casefold()] = name
    return list(names.values())


def sync(
    bitable: BitableClient, settings, parsed: ParsedTemplate, *, tz: tzinfo, apply: bool
) -> Summary:
    """把模板同步进 Base。``apply=False`` 只算不写。"""
    summary = Summary()

    # ---- 渠道 ----
    referral_primary = bitable.resolve_primary_field(settings.table_referral).name
    existing_referrals: dict[str, Record] = {}
    for record in bitable.iter_records(settings.table_referral):
        code = _clean_code(record.fields.get(schema.REFERRAL_NO))
        if code:
            existing_referrals.setdefault(code, record)

    to_create: list[dict[str, Any]] = []
    for row in parsed.referrals:
        desired = _referral_fields(row, referral_primary, tz)
        current = existing_referrals.get(row.code)
        if current is None:
            to_create.append(desired)
            continue
        changes = _changes(current.fields, desired)
        if changes:
            summary.referrals_updated += 1
            if apply:
                bitable.update_record(settings.table_referral, current.record_id, changes)
    summary.referrals_created = len(to_create)
    if apply and to_create:
        bitable.batch_create_records(settings.table_referral, to_create)
        existing_referrals = {}
        for record in bitable.iter_records(settings.table_referral):
            code = _clean_code(record.fields.get(schema.REFERRAL_NO))
            if code:
                existing_referrals.setdefault(code, record)
    referral_id_by_code = {code: record.record_id for code, record in existing_referrals.items()}
    template_codes = {row.code for row in parsed.referrals}

    # ---- 客户 ----
    client_primary = bitable.resolve_primary_field(settings.table_client).name
    code_by_referral_id = {record_id: code for code, record_id in referral_id_by_code.items()}
    by_uid: dict[str, Record] = {}
    by_key: dict[tuple[str, str], Record] = {}
    for record in bitable.iter_records(settings.table_client):
        uid = to_uid(record.fields.get(schema.CLIENT_UID))
        if uid:
            by_uid.setdefault(uid, record)
        linked = _link_ids(record.fields.get(schema.CLIENT_REFERRAL_LINK))
        code = code_by_referral_id.get(linked[0], "") if linked else ""
        by_key.setdefault((code, _norm_name(record.fields.get(schema.CLIENT_NAME))), record)

    client_creates: list[dict[str, Any]] = []
    for row in _dedupe_clients(parsed.clients, summary.notes):
        if row.code not in template_codes and row.code not in referral_id_by_code:
            summary.notes.append(f"客户「{row.name}」引用的渠道 {row.code} 不存在，没导")
            continue
        current = by_uid.get(row.uid) if row.uid else None
        if current is None:
            current = by_key.get((row.code, _norm_name(row.name)))
        referral_record_id = referral_id_by_code.get(row.code, "")
        if current is None:
            summary.clients_created += 1
            if apply:
                client_creates.append(_client_fields(row, referral_record_id, client_primary))
            continue
        changes = _changes(current.fields, _client_fields(row, referral_record_id, client_primary))
        if changes:
            summary.clients_updated += 1
            if apply:
                bitable.update_record(settings.table_client, current.record_id, changes)
    if apply and client_creates:
        bitable.batch_create_records(settings.table_client, client_creates)

    # ---- 销售名册：只补姓名，OpenID 等拿到了再填 ----
    existing_sales = {
        _clean_text(record.fields.get(schema.SALES_NAME)).casefold()
        for record in bitable.iter_records(settings.table_sales)
    }
    sales_creates = [
        {
            schema.SALES_NAME: name,
            schema.SALES_ROLE: schema.ROLE_SALES,
            schema.SALES_STATUS: schema.SALES_STATUS_ACTIVE,
        }
        for name in _sales_names(parsed)
        if name.casefold() not in existing_sales
    ]
    summary.sales_created = len(sales_creates)
    if apply and sales_creates:
        bitable.batch_create_records(settings.table_sales, sales_creates)

    logger.info(
        "渠道 新增 %d 更新 %d；客户 新增 %d 更新 %d；销售名册 新增 %d；%s",
        summary.referrals_created,
        summary.referrals_updated,
        summary.clients_created,
        summary.clients_updated,
        summary.sales_created,
        "已写入" if apply else "预演",
    )
    return summary


# ---------- 命令行 ----------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="把渠道登记表和客户表从模板 xlsx 导进 Base")
    parser.add_argument("--file", required=True, help="模板 xlsx 的路径")
    parser.add_argument("--apply", action="store_true", help="真的写，默认只预演")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")
    settings = load_settings()
    require_settings(
        settings, "LARK_BASE_APP_TOKEN", "TABLE_REFERRAL", "TABLE_CLIENT", "TABLE_SALES"
    )
    return run(args, settings, BitableClient(settings.base_app_token))


def run(args: argparse.Namespace, settings, bitable: BitableClient) -> int:
    try:
        parsed = parse_template(Path(args.file))
    except RegistrationImportError as exc:
        print(f"\n导入失败：{exc}", file=sys.stderr)
        return 1

    with_uid = sum(1 for c in parsed.clients if c.uid)
    print(f"读取：{args.file}")
    print(
        f"渠道 {len(parsed.referrals)} 个；客户 {len(parsed.clients)} 个，带 UID 的 {with_uid} 个"
    )
    for line in parsed.warnings:
        print(f"  ! {line}")

    tz = ZoneInfo(settings.business_timezone)
    summary = sync(bitable, settings, parsed, tz=tz, apply=args.apply)

    verb = "已" if args.apply else "将"
    print(
        f"\n渠道表：{verb}新增 {summary.referrals_created}，{verb}更新 {summary.referrals_updated}"
    )
    print(f"客户表：{verb}新增 {summary.clients_created}，{verb}更新 {summary.clients_updated}")
    print(f"销售名册：{verb}新增 {summary.sales_created}（OpenID 留空，拿到后自己填）")
    for line in summary.notes:
        print(f"  · {line}")
    if not args.apply:
        print("\n预演没写 Base。确认无误后加 --apply 真正执行。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
