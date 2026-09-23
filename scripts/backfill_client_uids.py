#!/usr/bin/env python
"""按客户名把「全量 UID」表里的 UID 补进客户表空着的那一列。

背景：一部分客户当初登记时没填 UID（客户表那一格是空的）。UID 是看板 join 客户表的
唯一钥匙 —— 缺了它，这个客户以后无论做多少交易都挂不上渠道，佣金一直漏。

**这个脚本按姓名匹配，而整套系统别的地方刻意都不用姓名**（见 verify_commission.py
开头：姓名会撞车、会大小写不一致、会因为先名后姓对不上）。匹配错的后果是这个客户的
收入算到别的渠道头上，钱付给错的人，而且报表上看不出任何异常 —— 有值、有比例、有
金额，一切正常。实测那份全量表里 ``HAI LIU`` 就对着两个不同 UID，光看名字无法判断
是哪一个。

所以这里的取舍是**宁可少补几条让人手工判断，也绝不猜**：

  · 名字在来源表里恰好命中一条        -> 补
  · 名字在来源表里有多条不同 UID      -> 不补，列出来
  · 名字在来源表里找不到              -> 不补，列出来（附一个最接近的名字当提示，不据此动手）
  · 配到的 UID 已被客户表另一条占用   -> 不补，列出来
  · 客户表两条同名空行会配到同一个 UID -> 都不补，列出来

默认只预演，加 ``--apply`` 才真写。

    uv run python scripts/backfill_client_uids.py --file 全量UID.xlsx
    uv run python scripts/backfill_client_uids.py --file 全量UID.xlsx --apply

⚠️ 补完之后必须重跑一次完整的看板导入，那些客户的历史交易行才会挂上「客户」关联：
    uv run python scripts/import_daily_board.py --file attachments/OTC组销售明细_YYYY-MM-DD.xlsx
"""

from __future__ import annotations

import argparse
import difflib
import logging
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

# 复用登记导入那套 UID 安全读法：浮点数直接报错、科学计数法和「尾巴一串 0」判损。
# 不在这里另写一份 —— 两处各写一份，改了一处忘了另一处就是一个静默错账。
import import_registrations as reg  # noqa: E402

from crm_basebot.domain import schema  # noqa: E402
from crm_basebot.lark.bitable import BitableClient  # noqa: E402
from crm_basebot.lark.values import (  # noqa: E402
    assess_uid_health,
    to_uid,
    uid_health_advice,
)
from crm_basebot.startup import load_settings, require_settings  # noqa: E402

logger = logging.getLogger(__name__)

# 来源表的表头 -> 内部键。飞书表格导出的那份用 user_id / client_name，
# 登记模板那份用 UID / 客户名称，两套都认。多出来的列（org_id、type、kyc_date）忽略。
SOURCE_COLUMNS = {
    "user_id": "uid",
    "userid": "uid",
    "uid": "uid",
    "客户uid": "uid",
    "client_name": "name",
    "client name": "name",
    "clientname": "name",
    "客户名称": "name",
    "name": "name",
}

# 损坏比例超过这个数就不让写了：说明整份导出是按数字处理的，不只是个别单元格的问题。
DAMAGED_RATIO_LIMIT = 0.2


def parse_source(
    path: Path, sheet_name: str | None
) -> tuple[dict[str, set[str]], list[str], list[str]]:
    """来源 xlsx -> ({规整后的客户名: {UID, ...}}, 采用的 UID, 警告)。

    同名的多个 UID **全部留下**，不像模板导入那样只取第一个 —— 这里要靠「有几个」
    判断能不能补，提前去重就把歧义藏起来了。

    **不能用 openpyxl 的 read_only 模式。** 飞书表格导出的 xlsx 里
    ``<dimension ref="A1">`` 是错的（实测：1032 行的表只声明了一格），read_only 模式
    信这个声明，结果只读到表头就停了，看起来像「文件是空的」。普通模式不信它，自己
    扫完整张表。
    """
    if not path.exists():
        raise reg.RegistrationImportError(f"找不到文件：{path}")

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)

    warnings: list[str] = []
    mapping: dict[str, set[str]] = defaultdict(set)
    accepted: list[str] = []
    damaged = 0
    try:
        sheet = (
            reg._find_sheet(workbook, sheet_name)
            if sheet_name
            else workbook[workbook.sheetnames[0]]
        )
        rows = reg._rows(sheet)
        if len(rows) < 2:
            raise reg.RegistrationImportError(f"工作表「{sheet.title}」没有数据行")
        col = reg._columns(rows[0], SOURCE_COLUMNS, ("uid", "name"), sheet=sheet.title)

        for row_num, row in enumerate(rows[1:], start=2):
            raw_name = row[col["name"]] if col["name"] < len(row) else None
            raw_uid = row[col["uid"]] if col["uid"] < len(row) else None
            name = reg._norm_name(raw_name)
            if not name or raw_uid in (None, ""):
                continue
            uid = reg._cell_to_uid(raw_uid, where=f"「{sheet.title}」第 {row_num} 行")
            reason = reg._damaged_uid_reason(uid)
            if reason:
                damaged += 1
                warnings.append(f"第 {row_num} 行「{raw_name}」的 UID {uid} {reason}，没有采用")
                continue
            mapping[name].add(uid)
            accepted.append(uid)
    finally:
        workbook.close()

    total = len(accepted) + damaged
    if total and damaged / total > DAMAGED_RATIO_LIMIT:
        raise reg.RegistrationImportError(
            f"来源表 {total} 个 UID 里有 {damaged} 个已经被 Excel 改坏"
            f"（超过 {DAMAGED_RATIO_LIMIT:.0%}）—— 整份导出像是按数字处理的，不能拿来补。"
            "请让对方把 UID 那一列设成【文本】后重新导出。"
        )

    return dict(mapping), accepted, warnings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按客户名补客户表里空着的 UID")
    parser.add_argument("--file", required=True, help="全量 UID 的 xlsx")
    parser.add_argument("--sheet", help="工作表名，默认第一张")
    parser.add_argument("--apply", action="store_true", help="真写；不加则只预演")
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")

    settings = load_settings()
    require_settings(settings, "LARK_BASE_APP_TOKEN", "TABLE_CLIENT")
    bitable = BitableClient(settings.base_app_token)

    # ---------- 读来源 ----------
    source, accepted, warnings = parse_source(Path(args.file), args.sheet)
    print(f"来源表：{len(source)} 个客户名，采用 {len(accepted)} 个 UID")
    for line in warnings:
        print(f"  ⚠️ {line}")
    print(f"  UID 体检：{uid_health_advice(assess_uid_health(accepted))}")

    ambiguous = {name for name, uids in source.items() if len(uids) > 1}
    if ambiguous:
        print(f"  来源表里有 {len(ambiguous)} 个名字对着多个 UID，这些名字一律不补")

    # ---------- 读客户表 ----------
    # 已经有 UID 的行：它们占用的 UID 不能再被别人配上。重复 UID 会让看板的匹配变成
    # 不确定 —— 导入脚本按 UID 找客户时只认第一条。
    taken: dict[str, str] = {}
    blanks: list[tuple[str, str]] = []  # (record_id, 客户名称原文)
    total = 0
    for record in bitable.iter_records(
        settings.table_client, field_names=[schema.CLIENT_UID, schema.CLIENT_NAME]
    ):
        total += 1
        name = reg._clean_text(record.fields.get(schema.CLIENT_NAME))
        uid = to_uid(record.fields.get(schema.CLIENT_UID))
        if uid:
            taken.setdefault(uid, name or record.record_id)
        else:
            blanks.append((record.record_id, name))

    print(f"\n客户表：{total} 条，其中 {len(blanks)} 条没有 UID")
    if not blanks:
        print("没有要补的。")
        return 0

    # ---------- 配对 ----------
    source_names = list(source)
    planned: list[tuple[str, str, str]] = []  # (record_id, 客户名, uid)
    skipped: list[tuple[str, str]] = []  # (客户名, 原因)
    claimed: dict[str, list[str]] = defaultdict(list)  # uid -> 想要它的客户名

    for record_id, name in blanks:
        key = reg._norm_name(name)
        if not key:
            skipped.append((f"(记录 {record_id})", "这条记录连客户名称都是空的"))
            continue

        uids = source.get(key)
        if not uids:
            hint = difflib.get_close_matches(key, source_names, n=1, cutoff=0.85)
            reason = "来源表里找不到这个名字"
            if hint:
                reason += f"（最接近的是「{hint[0]}」—— 只是提示，脚本不据此动手）"
            skipped.append((name, reason))
            continue

        if len(uids) > 1:
            skipped.append((name, f"来源表里这个名字对着 {len(uids)} 个不同 UID：{sorted(uids)}"))
            continue

        uid = next(iter(uids))
        if uid in taken:
            skipped.append((name, f"UID {uid} 已经被客户表里的「{taken[uid]}」用了"))
            continue

        claimed[uid].append(name)
        planned.append((record_id, name, uid))

    # 客户表里两条同名空行会配到同一个 UID —— 两条都不补，先让人把重名处理掉。
    conflicted = {uid for uid, names in claimed.items() if len(names) > 1}
    for uid in sorted(conflicted):
        skipped.append(
            (
                "、".join(claimed[uid]),
                f"客户表里 {len(claimed[uid])} 条同名空行都会配到 UID {uid}，都没补",
            )
        )
    planned = [item for item in planned if item[2] not in conflicted]

    # ---------- 输出 ----------
    if skipped:
        print(f"\n没补的 {len(skipped)} 条：")
        for name, reason in skipped:
            print(f"  {name:<44} {reason}")

    if not planned:
        print("\n没有可以安全补上的记录。")
        return 0

    print(f"\n可以补的 {len(planned)} 条：")
    for _, name, uid in planned:
        print(f"  {name:<44} -> {uid}")

    if not args.apply:
        print("\n预演：没有写 Base。确认无误后加 --apply。")
        return 0

    # ---------- 写 ----------
    # 一条一条写：Bitable 的写接口不支持并发（1254291 Write conflict）。
    # 每条写完回读确认 —— 写接口返回成功不等于值进去了（类型不符会被静默丢弃）。
    written = 0
    for record_id, name, uid in planned:
        bitable.update_record(settings.table_client, record_id, {schema.CLIENT_UID: uid})
        after = to_uid(
            bitable.get_record(settings.table_client, record_id).fields.get(schema.CLIENT_UID)
        )
        if after != uid:
            print(
                f"\n「{name}」写回读到的值不对（{after!r}，期望 {uid}）—— 停下来。"
                "请检查客户表 UID 那一列是不是被改成了数字类型。",
                file=sys.stderr,
            )
            return 1
        written += 1
        print(f"  ✓ {name} -> {uid}")

    print(f"\n写入并回读确认 {written} 条。")
    print(
        "\n下一步（必须做，否则这些客户的历史交易行还是挂不上渠道）："
        "\n  uv run python scripts/import_daily_board.py --file attachments/<最新那份>.xlsx"
        "\n再跑 uv run python scripts/verify_commission.py 核对，「漏挂」应该还是 0。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
