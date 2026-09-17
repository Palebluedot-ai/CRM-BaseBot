#!/usr/bin/env python
"""把内部系统导出的交易明细 xlsx 导入 Base 的 Daily Revenue Board 表。

## 表头就是合同

xlsx 的表头以 2026-09-17 的「OTC组销售明细」导出为准：18 列逐字对应
``schema.DAILY_BOARD_FIELDS``，Base 里的列名和它一模一样，不做改名。导出里少任何一列，
导入直接拒绝并点名缺的是哪列；多出来的列忽略并提示。导出工具偶尔把括号写成全角、
表头前后带空格，这类差异规整后再比，不因为一个括号的宽度拒掉整份文件。

## 只导新加坡站

导出里有香港站、新加坡站、中东站三个站点，看板只要新加坡站的记录（2026-09-17 定的）。
其他站点的行解析完就丢掉，不进 Base。筛的是「站点」列，不是「销售分组」列：新加坡站
的记录里也有 HK组、支付组的销售。

先删后写替换的是导出覆盖到的**所有**日期，包括那天只有其他站点记录的日子：导出说这天
新加坡站没有记录，Base 里这天的旧记录就该清掉。筛完一行新加坡站都不剩时拒绝导入，
否则会把这些日期的记录删光。

## 为什么必须走 xlsx 不走 CSV

用户ID 大多是 18-19 位数字。CSV 在源头几乎总是过一手 Excel，而 Excel 只保留 15 位
有效数字，低位在存成 CSV 之前就被抹成 0 了，下游修不了。xlsx 保留单元格的原生类型，
只要用户ID 那一列在源头是文本，openpyxl 拿到的就是原值。这个脚本对用户ID 强制字符串
校验，拿到浮点数直接报错而不是凑合。

## 幂等：按交易日期先删后写

导出每天更新，脚本会被反复跑。同一用户同一天可能有多行（那份导出里有 402 组），
没有稳定的行主键，按行 upsert 会漏改。所以按日期整批替换：

    · 默认：xlsx 里出现的每一个交易日期，先删 Base 上这些日期的记录，再写新的
    · --date YYYY-MM-DD：只处理这一天
    · --dry-run：只解析和统计，不碰 Base

## 批量写

一行一个请求的话，1.2 万行就是 1.2 万次调用，免费版一个月的基线额度才 1 万次。
删和写都按每批 500 条发，1.2 万行大约 25 次写入。

## 日期按业务时区

「交易日期」「KYC日期」是新加坡的日历日。写进 Base 时取 BUSINESS_TIMEZONE 那一天的
零点，先删后写也按这个时区取日期，界面里看到的就是那一天 0:00。

## 用法

    uv run python scripts/import_daily_board.py --file 明细.xlsx --dry-run   # 只解析不写
    uv run python scripts/import_daily_board.py --file 明细.xlsx             # 导入
    uv run python scripts/import_daily_board.py --file 明细.xlsx --date 2026-09-16
    uv run python scripts/import_daily_board.py        # 用 .env 里的 DAILY_BOARD_XLSX
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from collections import Counter, defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, tzinfo
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from crm_basebot.domain import schema  # noqa: E402
from crm_basebot.domain.dates import date_to_ms, ms_to_date  # noqa: E402
from crm_basebot.lark.bitable import (  # noqa: E402
    FIELD_TYPE_DATETIME,
    FIELD_TYPE_NUMBER,
    MAX_BATCH_SIZE,
    BitableClient,
)
from crm_basebot.lark.values import PrecisionLossError, extract_text, to_uid  # noqa: E402
from crm_basebot.startup import load_settings, require_settings  # noqa: E402

logger = logging.getLogger(__name__)

# xlsx 必须有的表头，也就是 Base 里看板表的列。
EXPECTED_HEADERS: tuple[str, ...] = tuple(schema.DAILY_BOARD_FIELDS)

DATE_FIELDS = frozenset(
    name for name, kind in schema.DAILY_BOARD_FIELDS.items() if kind == FIELD_TYPE_DATETIME
)
NUMBER_FIELDS = frozenset(
    name for name, kind in schema.DAILY_BOARD_FIELDS.items() if kind == FIELD_TYPE_NUMBER
)

# 一行里缺了这几个值就没法算佣金，跳过这一行。其余列都允许为空。
REQUIRED_VALUES: tuple[str, ...] = (
    schema.BOARD_CLIENT_UID,
    schema.BOARD_ORDER_DATE,
    schema.BOARD_TOTAL_REVENUE,
)

# 导出工具偶尔写成全角的符号，规整成 schema 里用的半角再比。
_HEADER_TRANSLATION = str.maketrans({"（": "(", "）": ")", "＿": "_", "＋": "+"})
_WHITESPACE = re.compile(r"\s+")


def normalize_header(raw: Any) -> str:
    """表头规整：去掉所有空白，全角括号、下划线、加号换成半角。"""
    if raw is None:
        return ""
    return _WHITESPACE.sub("", str(raw)).translate(_HEADER_TRANSLATION)


@dataclass(frozen=True)
class BoardRow:
    order_date: date
    fields: dict[str, Any]


class BoardImportError(RuntimeError):
    """xlsx 结构或值不符合要求。"""


def _load_workbook(path: Path):
    """惰性 import openpyxl，避免让主进程启动就吃这个依赖。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BoardImportError("缺 openpyxl，先跑 `uv sync` 装齐依赖再试。") from exc

    if not path.exists():
        raise BoardImportError(f"找不到文件：{path}")
    # data_only=True 让公式单元格返回上次计算的值，而不是公式本身
    return load_workbook(path, data_only=True, read_only=True)


def _read_header(sheet) -> list[str]:
    """读第一行表头，返回每一列对应的看板列名，认不出的列是空串。

    缺任何一列都报错：少了哪列，下游就会拿着空值往下算，宁可在导入前停下。
    同一列出现两次也报错：不知道该用哪一列。
    """
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if first_row is None:
        raise BoardImportError("xlsx 是空的，连表头都没有")

    raw_headers = ["" if cell is None else str(cell) for cell in first_row]
    expected = set(EXPECTED_HEADERS)
    normalized = [normalize_header(raw) for raw in raw_headers]
    columns = [name if name in expected else "" for name in normalized]

    missing = [name for name in EXPECTED_HEADERS if name not in columns]
    if missing:
        raise BoardImportError(
            f"xlsx 表头缺少这些列：{'、'.join(missing)}\n"
            f"实际表头：{[raw for raw in raw_headers if raw]}\n"
            "表头以 2026-09-17 的「OTC组销售明细」导出为准，清单见 schema.DAILY_BOARD_FIELDS。"
        )

    duplicated = sorted({name for name in columns if name and columns.count(name) > 1})
    if duplicated:
        raise BoardImportError(f"xlsx 表头里这些列出现了不止一次：{'、'.join(duplicated)}")

    unknown = [
        raw for raw, name in zip(raw_headers, columns, strict=True) if raw.strip() and not name
    ]
    if unknown:
        logger.warning("忽略看板表里没有的列：%s", unknown)

    return columns


def _cell_to_uid(value: Any, *, row_num: int) -> str:
    """用户ID 的值。任何浮点形态都当作精度已损伤，直接拒绝。

    xlsx 里用户ID 应该是文本；拿到 float 就是源头把它当数字存了，18-19 位的数字
    过一遍 float 就丢低位，下游对不上。让这里就炸，而不是继续往下算。
    """
    column = schema.BOARD_CLIENT_UID
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        raise BoardImportError(f"第 {row_num} 行的「{column}」是布尔值")
    if isinstance(value, float):
        raise BoardImportError(
            f"第 {row_num} 行的「{column}」是浮点数，说明源头把这一列当数字存了。"
            f"请让上游把「{column}」那一列在 xlsx 里设成文本格式重新导出，"
            "否则 18-19 位的 ID 低位已经被抹成 0。"
        )
    if isinstance(value, int):
        return str(value)
    try:
        return to_uid(value)
    except PrecisionLossError as exc:
        raise BoardImportError(f"第 {row_num} 行的「{column}」无法安全转成字符串：{exc}") from exc


def _cell_to_date(value: Any, column: str, *, row_num: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        # 只取日期部分（扔掉可能跟着的时分秒），再规范化分隔符
        text = value.strip().split(" ", 1)[0].replace("/", "-")
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            pass
    raise BoardImportError(f"第 {row_num} 行的「{column}」无法解析成日期：{value!r}")


def _cell_to_number(value: Any, column: str, *, row_num: int) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):  # bool 是 int 的子类，得先挡
        raise BoardImportError(f"第 {row_num} 行的「{column}」是布尔值 {value!r}")
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, str):
        stripped = value.strip().replace(",", "")
        if not stripped:
            return None
        try:
            return float(stripped)
        except ValueError as exc:
            raise BoardImportError(f"第 {row_num} 行的「{column}」不是数字：{value!r}") from exc
    raise BoardImportError(f"第 {row_num} 行的「{column}」无法转数字：{value!r}")


def _cell_value(cell: Any, column: str, *, row_num: int) -> Any:
    """按列的类型转换一个单元格。返回 None 表示这个单元格是空的，不写这一列。"""
    if column == schema.BOARD_CLIENT_UID:
        return _cell_to_uid(cell, row_num=row_num) or None
    if column in DATE_FIELDS:
        return _cell_to_date(cell, column, row_num=row_num)
    if column in NUMBER_FIELDS:
        return _cell_to_number(cell, column, row_num=row_num)
    return extract_text(cell) or None


def parse_workbook(path: Path, *, only_date: date | None = None) -> list[BoardRow]:
    """把 xlsx 读成 BoardRow 列表。所有校验都在这里做，不碰 Base。

    ``only_date`` 传了就只保留那一天的行。
    """
    workbook = _load_workbook(path)
    rows: list[BoardRow] = []
    skipped = 0
    try:
        sheet = workbook.active
        columns = _read_header(sheet)

        for index, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None or cell == "" for cell in raw_row):
                continue  # 空行

            fields: dict[str, Any] = {}
            for cell, column in zip(raw_row, columns, strict=False):
                if not column or cell is None or cell == "":
                    continue
                value = _cell_value(cell, column, row_num=index)
                if value is not None:
                    fields[column] = value

            # 缺值按「这一列有没有值」判断，不按真假判断：总收入是 0 的行照样要导。
            missing = [column for column in REQUIRED_VALUES if column not in fields]
            if missing:
                # 只报行号和缺了哪列。日志会被转发和截图，客户名称、金额这些不打出来。
                labels = "、".join(f"「{column}」" for column in missing)
                logger.warning("第 %d 行缺少%s，跳过", index, labels)
                skipped += 1
                continue

            order_date: date = fields[schema.BOARD_ORDER_DATE]
            if only_date is not None and order_date != only_date:
                continue

            rows.append(BoardRow(order_date=order_date, fields=fields))
    finally:
        workbook.close()

    if skipped:
        logger.warning("一共跳过 %d 行缺少必填值的记录", skipped)
    return rows


def _existing_by_date(
    bitable: BitableClient, table_id: str, *, tz: tzinfo
) -> dict[date, list[str]]:
    """扫一遍 Base 表，按交易日期分组 record_id，供「先删」阶段用。

    日期按业务时区取：界面里手工填的「9 月 10 日」是新加坡零点，按 UTC 取会成 9 月 9 日，
    重导 9 月 10 日时就删不掉它。
    """
    grouped: dict[date, list[str]] = defaultdict(list)
    for record in bitable.iter_records(table_id, field_names=[schema.BOARD_ORDER_DATE]):
        raw = record.fields.get(schema.BOARD_ORDER_DATE)
        if isinstance(raw, int | float) and not isinstance(raw, bool):
            grouped[ms_to_date(raw, tz=tz)].append(record.record_id)
    return grouped


def _to_payload(row: BoardRow, *, tz: tzinfo) -> dict[str, Any]:
    """一行写进 Base 的字段：日期列换成业务时区那天零点的毫秒时间戳，其余原样。"""
    return {
        column: date_to_ms(value, tz=tz) if column in DATE_FIELDS else value
        for column, value in row.fields.items()
    }


def _apply(
    bitable: BitableClient,
    table_id: str,
    rows: list[BoardRow],
    *,
    tz: tzinfo,
    replace_dates: Iterable[date] | None = None,
) -> tuple[int, int]:
    """先删要替换的日期上的旧记录，再写新记录，都按批发。返回 (删除数, 写入数)。

    ``replace_dates`` 是要整天替换的日期，不传就取 rows 覆盖到的日期。
    导入时传导出覆盖到的全部日期：某天导出里只有其他站点的行，这天的旧记录也要删掉。
    ``tz`` 是业务时区：日期列写成那一天在业务时区的零点，删旧行也按同一时区取日期。
    """
    if replace_dates is None:
        replace_dates = {row.order_date for row in rows}
    days = sorted(set(replace_dates))
    existing = _existing_by_date(bitable, table_id, tz=tz)
    stale = [record_id for day in days for record_id in existing.get(day, [])]

    deleted = bitable.batch_delete_records(table_id, stale)
    written = bitable.batch_create_records(table_id, [_to_payload(row, tz=tz) for row in rows])
    return deleted, written


def split_by_station(rows: list[BoardRow]) -> tuple[list[BoardRow], Counter[str]]:
    """只留看板要的站点，顺带数出每个站点各有多少行，方便核对筛得对不对。"""
    counts: Counter[str] = Counter(
        row.fields.get(schema.BOARD_STATION, "（站点为空）") for row in rows
    )
    kept = [
        row for row in rows if row.fields.get(schema.BOARD_STATION) == schema.BOARD_STATION_IN_SCOPE
    ]
    return kept, counts


def _describe_stations(counts: Counter[str]) -> str:
    in_scope = schema.BOARD_STATION_IN_SCOPE
    text = f"{in_scope} {counts.get(in_scope, 0)} 行导入"
    others = [f"{station} {n} 行" for station, n in counts.most_common() if station != in_scope]
    if others:
        text += f"；{'、'.join(others)}不导入"
    return text


def _print_summary(rows: list[BoardRow]) -> None:
    """按月列行数和天数。几个月的数据按天列会刷出几百行，看不出东西。"""
    rows_by_month: dict[str, int] = defaultdict(int)
    days_by_month: dict[str, set[date]] = defaultdict(set)
    for row in rows:
        month = row.order_date.strftime("%Y-%m")
        rows_by_month[month] += 1
        days_by_month[month].add(row.order_date)

    days = {row.order_date for row in rows}
    print(f"\n{len(rows)} 行，{min(days)} 到 {max(days)}，覆盖 {len(days)} 天：")
    for month in sorted(rows_by_month):
        print(f"  {month}  {rows_by_month[month]:>6} 行  {len(days_by_month[month]):>3} 天")

    batches = -(-len(rows) // MAX_BATCH_SIZE)
    print(f"写入每批 {MAX_BATCH_SIZE} 条，约 {batches} 次调用；删除旧记录另算。")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="把交易明细 xlsx 里新加坡站的记录导入 Base 的日读看板表"
    )
    parser.add_argument("--file", help="xlsx 文件路径。不传就用 .env 里的 DAILY_BOARD_XLSX")
    parser.add_argument("--date", help="只导这一天 YYYY-MM-DD。xlsx 里其他日期的行忽略")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只解析和统计，不碰 Base。先用它核对表头和数据",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")

    settings = load_settings()
    require_settings(settings, "LARK_BASE_APP_TOKEN", "TABLE_DAILY_BOARD")
    return run(args, settings, BitableClient(settings.base_app_token))


def run(args: argparse.Namespace, settings, bitable: BitableClient) -> int:
    """入口的主体。settings 和 bitable 从外面传进来，测试里换成假件就能把整条路走一遍。"""
    xlsx_path = Path(args.file or settings.daily_board_xlsx or "")
    if not xlsx_path or str(xlsx_path) == ".":
        print(
            "没指定 xlsx：加 --file /path/to.xlsx，或在 .env 里设 DAILY_BOARD_XLSX。",
            file=sys.stderr,
        )
        return 1

    only_date: date | None = None
    if args.date:
        try:
            only_date = datetime.strptime(args.date, "%Y-%m-%d").date()
        except ValueError:
            print(f"--date 格式要是 YYYY-MM-DD，你给的是 {args.date!r}", file=sys.stderr)
            return 1

    print(f"读取：{xlsx_path}")
    if only_date:
        print(f"只导：{only_date}")

    try:
        rows = parse_workbook(xlsx_path, only_date=only_date)
    except BoardImportError as exc:
        print(f"\n导入失败：{exc}", file=sys.stderr)
        return 1

    if not rows:
        print("没有匹配的行可导入。")
        return 0

    kept, counts = split_by_station(rows)
    print(f"站点：{_describe_stations(counts)}")
    if not kept:
        found = "、".join(f"{station} {n} 行" for station, n in counts.most_common())
        print(
            "\n没有导入：这份导出里一行"
            f"「{schema.BOARD_STATION_IN_SCOPE}」都没有，只有 {found}。\n"
            "Base 没有动。看一眼是不是导错了文件，或者站点的写法变了。",
            file=sys.stderr,
        )
        return 1

    _print_summary(kept)

    if args.dry_run:
        print("\n--dry-run：只解析没写 Base。去掉这个开关才会真正导入。")
        return 0

    tz = ZoneInfo(settings.business_timezone)
    print(
        "\n先删导出覆盖到的日期上的旧记录，再写新记录。"
        f"日期按 {settings.business_timezone} 的零点写入。"
    )
    covered = {row.order_date for row in rows}
    deleted, written = _apply(
        bitable, settings.table_daily_board, kept, tz=tz, replace_dates=covered
    )
    print(f"删除 {deleted} 条，写入 {written} 条。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
