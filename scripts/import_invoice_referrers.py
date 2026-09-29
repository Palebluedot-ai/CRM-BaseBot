#!/usr/bin/env python
"""把 invoice 小工具里存的收款资料，一次搬进 Base 的「Referral Information」。

    uv run python scripts/import_invoice_referrers.py            # 预演：列出会补哪些渠道的哪几项
    uv run python scripts/import_invoice_referrers.py --apply    # 真写
    uv run python scripts/import_invoice_referrers.py --xlsx ~/Downloads/Data.xlsx  # 老表

2026-09-29 定的：收款资料（地址、银行账户、钱包地址）以后只在 Base 维护一份，机器人的
「登记收款资料」改它、「生成 Invoice」读它。以前这些存在 invoice 小工具
（HashKey OTC Invoice Generator）自己的数据库里，这个脚本把那一份搬过来。

**要在装着 invoice 小工具的那台 Mac 上跑**（数据库在那台机器上）。默认读
``~/Library/Application Support/HashKey OTC Invoice Generator/app.db``，别的位置用 ``--db``。
收款资料只在这台电脑和飞书之间走，不经过任何别的地方。

## 怎么对上

**按名字对，不按编号。** invoice 小工具的交接说明写过：同一个编号在不同来源里会对到不同的人，
编号不可信。先比规整后的名字（大小写、空白、全半角都不算差别），对不上再比「词序无关」
（``KE JIAHUI`` 和 ``JIAHUI KE`` 算同一个）。一个名字对上 Base 里两个渠道的，跳过，列出来让人定。

## 绝不做的事

**Base 里已经有值的格子一律不改。** 只补空的。两边不一样的列出来让人看 —— 收款账号改错
一位，钱就打到别人那里去了。

输出里**不打印账号和钱包地址**，只说补了哪几项：这段输出可能被截图发出去。
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from crm_basebot.domain import schema  # noqa: E402
from crm_basebot.domain.names import norm, tokens  # noqa: E402
from crm_basebot.lark.bitable import BitableClient  # noqa: E402
from crm_basebot.lark.values import extract_text  # noqa: E402
from crm_basebot.startup import load_settings, require_settings  # noqa: E402

DEFAULT_DB = (
    Path.home() / "Library" / "Application Support" / "HashKey OTC Invoice Generator" / "app.db"
)

# invoice 小工具的 payment_method -> Base 的「收款方式」
METHODS = {"USD": schema.PAY_METHOD_BANK, "CRYPTO": schema.PAY_METHOD_CRYPTO}

PAYMENT_COLUMNS = (
    schema.REFERRAL_PAY_METHOD,
    schema.REFERRAL_BANK_ACCOUNT_NAME,
    schema.REFERRAL_BANK_NAME,
    schema.REFERRAL_BANK_ACCOUNT_NO,
    schema.REFERRAL_CRYPTO_TYPE,
    schema.REFERRAL_WALLET,
)


@dataclass(frozen=True)
class Source:
    code: str
    name: str
    fields: dict[str, str]


def _fields(
    *,
    method_raw: str,
    crypto_type: str,
    wallet: str,
    bank_account_name: str,
    bank_name: str,
    bank_account_no: str,
    address_lines: list[str],
) -> dict[str, str]:
    """一个 referrer 的收款资料 -> 要写进 Base 的字段（只放有值的）。

    invoice 小工具把「没写收款方式」也记成 CRYPTO（它的默认值），这里不跟：没有币种也没有
    钱包地址的，收款方式留空，免得 Base 里出现一个「加密货币」却没有钱包的渠道。
    """
    raw = method_raw.strip().upper()
    if raw == "USD":
        method = schema.PAY_METHOD_BANK
    elif raw.startswith("USDT") or crypto_type or wallet:
        method = schema.PAY_METHOD_CRYPTO
    else:
        method = ""
    fields = {
        schema.REFERRAL_ADDRESS: "\n".join(line.strip() for line in address_lines if line.strip()),
        schema.REFERRAL_PAY_METHOD: method,
    }
    if method == schema.PAY_METHOD_BANK:
        fields[schema.REFERRAL_BANK_ACCOUNT_NAME] = bank_account_name
        fields[schema.REFERRAL_BANK_NAME] = bank_name
        fields[schema.REFERRAL_BANK_ACCOUNT_NO] = bank_account_no
    elif method == schema.PAY_METHOD_CRYPTO:
        fields[schema.REFERRAL_CRYPTO_TYPE] = crypto_type or (raw if raw.startswith("USDT") else "")
        fields[schema.REFERRAL_WALLET] = wallet
    return {k: v.strip() for k, v in fields.items() if v and v.strip()}


def read_invoice_db(path: Path) -> list[Source]:
    """invoice 小工具的 referrers 表。"""
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        rows = connection.execute("SELECT * FROM referrers").fetchall()
    finally:
        connection.close()

    out: list[Source] = []
    for row in rows:
        keys = row.keys()

        def value(column: str, row=row, keys=keys) -> str:
            return str(row[column] or "").strip() if column in keys else ""

        out.append(
            Source(
                code=value("referrer_code"),
                name=value("name"),
                fields=_fields(
                    method_raw=value("payment_method"),
                    crypto_type=value("crypto_type"),
                    wallet=value("wallet_address"),
                    bank_account_name=value("bank_account_name"),
                    bank_name=value("bank_name"),
                    bank_account_no=value("bank_account_no"),
                    address_lines=value("address").splitlines(),
                ),
            )
        )
    return out


def _cell(row: tuple, index: int) -> str:
    value = row[index] if len(row) > index else None
    return "" if value is None else str(value).strip()


def read_data_xlsx(path: Path) -> list[Source]:
    """老的 Data.xlsx（invoice 小工具的导入源）的 Sheet2。

    列的位置照 invoice 小工具 import_referrers.py（从 0 数）：1 编号、2 名称、9 收款方式
    （USD / USDT-TRC / USDT-ERC）、10 ERC 钱包、11 TRC 钱包、12 户名、13 银行名、14 账号、
    18 地址。名称那一格空着的行是上一个 referrer 的地址续行。没有表头行。
    """
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Sheet2" not in book.sheetnames:
        raise SystemExit(f"{path} 里没有 Sheet2（有的是：{'、'.join(book.sheetnames)}）")

    parsed: list[dict] = []
    for row in book["Sheet2"].iter_rows(values_only=True):
        name, address = _cell(row, 2), _cell(row, 18)
        if name:
            if not _cell(row, 1):
                continue
            parsed.append(
                {
                    "code": _cell(row, 1),
                    "name": name,
                    "method": _cell(row, 9),
                    "wallet": _cell(row, 10) or _cell(row, 11),
                    "bank_account_name": _cell(row, 12),
                    "bank_name": _cell(row, 13),
                    "bank_account_no": _cell(row, 14),
                    "address": [address] if address else [],
                }
            )
        elif parsed and address:
            parsed[-1]["address"].append(address)

    return [
        Source(
            code=item["code"],
            name=item["name"],
            fields=_fields(
                method_raw=item["method"],
                crypto_type=item["method"] if item["method"].upper().startswith("USDT") else "",
                wallet=item["wallet"],
                bank_account_name=item["bank_account_name"],
                bank_name=item["bank_name"],
                bank_account_no=item["bank_account_no"],
                address_lines=item["address"],
            ),
        )
        for item in parsed
    ]


@dataclass(frozen=True)
class Target:
    record_id: str
    no: str
    name: str
    fields: dict[str, str]


def match(source: Source, targets: list[Target]) -> tuple[list[Target], str]:
    """先比规整名，再比词序无关。返回 (候选, 怎么对上的)。"""
    exact = [t for t in targets if norm(t.name) == norm(source.name)]
    if exact:
        return exact, ""
    loose = [t for t in targets if tokens(t.name) == tokens(source.name)]
    return loose, "（名字词序不同，核对一下）" if loose else ""


def plan_fill(source: Source, target: Target) -> tuple[dict[str, str], list[str]]:
    """只补 Base 里空着的格子。返回 (要写的字段, Base 已有但不一样的列名)。"""
    fill: dict[str, str] = {}
    differs: list[str] = []
    for column, value in source.fields.items():
        current = target.fields.get(column, "").strip()
        if not current:
            fill[column] = value
        elif current != value:
            differs.append(column)
    return fill, differs


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="把 invoice 小工具的收款资料搬进 Base")
    parser.add_argument("--db", default=str(DEFAULT_DB), help="invoice 小工具的 app.db")
    parser.add_argument("--xlsx", help="改读老的 Data.xlsx（Sheet2），不读 app.db")
    parser.add_argument("--apply", action="store_true", help="真写；不加则只预演")
    args = parser.parse_args(argv)

    if args.xlsx:
        source_path = Path(args.xlsx).expanduser()
        if not source_path.is_file():
            print(f"找不到 {source_path}")
            return 1
    else:
        source_path = Path(args.db).expanduser()
        if not source_path.is_file():
            print(f"找不到 invoice 小工具的数据库：{source_path}")
            print("  在装着 invoice 小工具的那台 Mac 上跑；数据库放在别处就加 --db 指过去，")
            print("  或者用 --xlsx 读老的 Data.xlsx。")
            return 1

    settings = load_settings()
    require_settings(settings, "LARK_BASE_APP_TOKEN", "TABLE_REFERRAL")
    bitable = BitableClient(settings.base_app_token)

    columns = {f.name for f in bitable.list_fields(settings.table_referral)}
    lacking = [c for c in PAYMENT_COLUMNS if c not in columns]
    if lacking:
        print(
            f"渠道表还没有这几列：{'、'.join(lacking)}。"
            "先跑：uv run python scripts/sync_base.py --apply"
        )
        return 1

    targets = [
        Target(
            record_id=record.record_id,
            no=extract_text(record.fields.get(schema.REFERRAL_NO)).strip(),
            name=extract_text(record.fields.get(schema.REFERRAL_NAME)).strip(),
            fields={
                column: extract_text(record.fields.get(column)).strip()
                for column in (schema.REFERRAL_ADDRESS, *PAYMENT_COLUMNS)
            },
        )
        for record in bitable.iter_records(settings.table_referral)
    ]
    targets = [t for t in targets if t.name]

    sources = read_data_xlsx(source_path) if args.xlsx else read_invoice_db(source_path)
    print(f"invoice 小工具里有 {len(sources)} 个 referrer，Base 里有 {len(targets)} 个渠道。\n")

    todo: list[tuple[Target, dict[str, str]]] = []
    unmatched: list[str] = []
    ambiguous: list[str] = []
    for source in sources:
        if not source.fields:
            continue
        candidates, note = match(source, targets)
        if not candidates:
            unmatched.append(f"{source.name}（小工具编号 {source.code or '无'}）")
            continue
        if len(candidates) > 1:
            names = "、".join(f"{t.no} {t.name}" for t in candidates)
            ambiguous.append(f"{source.name} -> {names}")
            continue
        target = candidates[0]
        fill, differs = plan_fill(source, target)
        label = f"{target.no} {target.name}{note}"
        if fill:
            print(f"  + {label}：补 {'、'.join(fill)}")
            todo.append((target, fill))
        if differs:
            print(f"  ! {label}：Base 已有、和小工具不一样，没改：{'、'.join(differs)}")

    if unmatched:
        print(f"\nBase 里找不到的 {len(unmatched)} 个（可能还没登记成渠道，或名字写法差太多）：")
        for line in unmatched:
            print(f"    {line}")
    if ambiguous:
        print(f"\n一个名字对上好几个渠道、跳过的 {len(ambiguous)} 个：")
        for line in ambiguous:
            print(f"    {line}")

    print(f"\n要补 {len(todo)} 个渠道。")
    if not args.apply:
        print("（预演，没写。确认无误后加 --apply）")
        return 0

    if todo:
        bitable.batch_update_records(
            settings.table_referral, {target.record_id: fill for target, fill in todo}
        )
    print(f"写完：补了 {len(todo)} 个渠道。以后改收款资料用机器人的「登记收款资料」。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
